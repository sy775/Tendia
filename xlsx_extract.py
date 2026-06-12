"""
xlsx_extract.py
---------------
Extraction texte des fichiers Excel (.xlsx, .xlsm) contenus dans les DCE.

Les xlsx publics typiques contiennent :
  - BPU (Bordereau Prix Unitaires) : Désignation/Quantité/Unité/PU/Total par ligne
  - DPGF (Décomposition Prix Global Forfaitaire) : structure de coût par poste
  - Cadre de réponse financière / technique : format imposé pour le candidat
  - Listes de lots et volumes détaillés
  - DUME / DC en xlsx (à filtrer comme bruit en amont si besoin)

Le LLM consomme ce texte pour compléter les sections Volumes / Montant /
Modalités du brief, qui étaient souvent vides en PDF-only.

Format de sortie pour chaque feuille :
    ### Feuille: <nom>

    val1 | val2 | val3 | ...
    val1 | val2 | val3 | ...
    ...

Choix : pipe-delimited (efficace en tokens, lisible par LLM), une feuille
par bloc séparé par une ligne vide, en-tête `### Feuille: <nom>` pour que
le LLM identifie clairement les sections.

Interface stable :
    extraire(chemin) -> str          texte concaténé ou "" si KO
    caracteres(chemin) -> int        compte sans dépendance sur tendia_extraire
"""

from pathlib import Path
from typing import Optional, Union

try:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    _OPENPYXL_OK = True
except ImportError:
    _OPENPYXL_OK = False


# Limite : nb de lignes par feuille, pour éviter d'exploser sur des xlsx
# énormes (logs capteurs, exports SIG, etc.). 2000 lignes couvre les BPU/DPGF
# raisonnables ; au-delà, on tronque avec une note explicite.
MAX_LIGNES_PAR_FEUILLE = 2000

# Limite colonnes : sécurité contre les feuilles avec 16k colonnes vides
# (parfois Excel garde des artefacts). 50 couvre les BPU les plus détaillés.
MAX_COLONNES = 50


def _cellule_vers_str(val) -> str:
    """Représentation compacte d'une valeur de cellule pour le LLM."""
    if val is None:
        return ""
    # int directement (évite "1.0" pour les entiers)
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        # Évite la notation scientifique sur les petits nombres ;
        # rstrip pour ne pas écrire "10.0000" mais "10"
        return f"{val:.4f}".rstrip("0").rstrip(".") or "0"
    if isinstance(val, int):
        return str(val)
    return str(val).strip()


def _ligne_vers_str(row, sep: str = " | ") -> str:
    return sep.join(_cellule_vers_str(c) for c in row)


def _est_ligne_vide(row) -> bool:
    return all(c is None or _cellule_vers_str(c) == "" for c in row)


def extraire(chemin: Union[Path, str], verbose: bool = False) -> str:
    """
    Lit un .xlsx/.xlsm et retourne son contenu en texte concaténé,
    une feuille après l'autre, en format pipe-delimited.

    Returns:
        Texte du classeur, ou "" si lecture impossible / fichier vide.
    """
    if not _OPENPYXL_OK:
        if verbose:
            print(f"    ⚠ openpyxl absent, impossible de lire xlsx")
        return ""

    chemin = Path(chemin)
    if not chemin.is_file():
        return ""

    try:
        # data_only=True : pour les cellules à formule, on prend la valeur
        # calculée (si présente). read_only=True : ne charge pas tout en RAM.
        wb = load_workbook(chemin, data_only=True, read_only=True)
    except (InvalidFileException, KeyError, ValueError, OSError) as e:
        if verbose:
            print(f"    ⚠ Erreur lecture {chemin.name} : "
                  f"{type(e).__name__}: {e}")
        return ""
    except Exception as e:
        # Catch-all car openpyxl peut lever des erreurs exotiques sur
        # fichiers corrompus (zlib.error, BadZipFile, etc.)
        if verbose:
            print(f"    ⚠ Erreur inattendue {chemin.name} : "
                  f"{type(e).__name__}: {e}")
        return ""

    morceaux: list[str] = []

    for ws in wb.worksheets:
        try:
            nom = ws.title or "Feuille"
        except Exception:
            nom = "Feuille"

        # Lecture ligne par ligne (mode read-only ne supporte pas l'accès
        # aléatoire ws[A1])
        lignes_brutes: list[list] = []
        tronque = False
        try:
            for row in ws.iter_rows(values_only=True):
                if len(lignes_brutes) >= MAX_LIGNES_PAR_FEUILLE:
                    tronque = True
                    break
                # Tronque les colonnes au-delà de MAX_COLONNES
                row_court = (row[:MAX_COLONNES]
                             if len(row) > MAX_COLONNES else row)
                lignes_brutes.append(list(row_court))
        except Exception as e:
            if verbose:
                print(f"    ⚠ Erreur itération feuille '{nom}' : "
                      f"{type(e).__name__}: {e}")
            continue

        if not lignes_brutes:
            continue

        # Tronque les colonnes vides à droite : calcule l'index max où une
        # cellule non vide existe sur l'ensemble des lignes.
        idx_max = -1
        for r in lignes_brutes:
            for i, c in enumerate(r):
                if c is not None and _cellule_vers_str(c) != "":
                    if i > idx_max:
                        idx_max = i
        if idx_max < 0:
            # Feuille entièrement vide
            continue
        lignes_courtes = [r[:idx_max + 1] for r in lignes_brutes]

        # Trim les lignes vides en début et fin
        debut, fin = 0, len(lignes_courtes)
        while debut < fin and _est_ligne_vide(lignes_courtes[debut]):
            debut += 1
        while fin > debut and _est_ligne_vide(lignes_courtes[fin - 1]):
            fin -= 1
        lignes_utiles = lignes_courtes[debut:fin]
        if not lignes_utiles:
            continue

        morceau = [f"### Feuille: {nom}", ""]
        for ligne in lignes_utiles:
            morceau.append(_ligne_vers_str(ligne))
        if tronque:
            morceau.append(
                f"... (feuille tronquée à {MAX_LIGNES_PAR_FEUILLE} lignes)"
            )
        morceau.append("")
        morceaux.append("\n".join(morceau))

    try:
        wb.close()
    except Exception:
        pass

    if not morceaux:
        return ""

    return "\n".join(morceaux)


def caracteres(chemin: Union[Path, str]) -> int:
    """Longueur du texte extrait. Utile pour le filtre 'fichier trop petit'
    si le pipeline tendia_extraire applique ce filtre aussi aux xlsx."""
    return len(extraire(chemin))


# =====================================================================
# CLI standalone : utile pour debug avant le branchement dans le pipeline
# =====================================================================

if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage: python xlsx_extract.py <fichier.xlsx> [...]")
        print("       Affiche le texte extrait, classeur par classeur.")
        sys.exit(1)

    for arg in sys.argv[1:]:
        chemin = Path(arg)
        if not chemin.is_file():
            print(f"❌ {arg} : fichier introuvable")
            continue
        txt = extraire(chemin, verbose=True)
        print(f"\n{'=' * 78}")
        print(f"  {chemin.name}  ({len(txt)} chars)")
        print(f"{'=' * 78}")
        if not txt:
            print("(vide ou non lisible)")
        else:
            print(txt)
