"""
pdf_extract.py
--------------
Extraction de texte depuis les fichiers déposés dans data/dce/<id_marche>/.

Supporte :
- .pdf  via pypdf (Règlement de consultation, CCFT, contrat, charte…)
- .docx via python-docx (Bordereau de Prix, cadre de réponse…)
- .xlsx via openpyxl (Base administrative, BPU au format Excel…)

L'inclusion .docx est cruciale : le BPU (Bordereau des Prix Unitaires) est
presque toujours en .docx ou .xlsx — sans ça, le LLM ne voit pas les prix.

Dépendances optionnelles (installer si tu rencontres l'erreur correspondante) :
    pip install pypdf            # .pdf
    pip install python-docx      # .docx
    pip install openpyxl         # .xlsx

API :
    extraire_fichier(chemin) -> str | None     # détecte le format automatiquement
    extraire_pdf(chemin) -> str | None         # PDF spécifique (rétro-compat)
    extraire_docx(chemin) -> str | None        # DOCX (texte + tableaux)
    extraire_xlsx(chemin) -> str | None        # XLSX (toutes feuilles)
    extraire_dossier(dossier) -> dict          # {nom_fichier: {hash, texte, taille_chars}}
    hash_fichier(chemin) -> str                # MD5 pour tracking
"""

import hashlib
import os
import re
from typing import Optional

# Imports avec fallback gracieux (pour ne pas casser si une lib manque)
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    import openpyxl
except ImportError:
    openpyxl = None


# Formats supportés (pour le scan de dossier)
EXTENSIONS_SUPPORTEES = (".pdf", ".docx", ".xlsx")


# ---------------------------------------------------------------------------
# Vérifications dépendances
# ---------------------------------------------------------------------------

def _check_pypdf():
    if PdfReader is None:
        raise RuntimeError(
            "pypdf n'est pas installé. Lance :\n"
            "    pip install pypdf"
        )


def _docx_disponible() -> bool:
    return DocxDocument is not None


def _xlsx_disponible() -> bool:
    return openpyxl is not None


# ---------------------------------------------------------------------------
# Hash de fichier (pour tracking)
# ---------------------------------------------------------------------------

def hash_fichier(chemin: str) -> str:
    """Retourne le hash MD5 d'un fichier (rapide, juste pour identifier)."""
    h = hashlib.md5()
    with open(chemin, "rb") as f:
        for bloc in iter(lambda: f.read(65536), b""):
            h.update(bloc)
    return h.hexdigest()


# ---------------------------------------------------------------------------
# Nettoyage texte commun
# ---------------------------------------------------------------------------

def _nettoyer_texte(texte: str) -> str:
    """Nettoyage léger du texte extrait."""
    # Joindre les lignes coupées par tiret en fin de ligne ("forma-\ntion")
    texte = re.sub(r"-\n(\w)", r"\1", texte)
    # Réduire les multiples sauts de ligne à 2 max
    texte = re.sub(r"\n{3,}", "\n\n", texte)
    # Réduire les multiples espaces à 1
    texte = re.sub(r"[ \t]+", " ", texte)
    return texte.strip()


# ---------------------------------------------------------------------------
# Extraction PDF
# ---------------------------------------------------------------------------

def extraire_pdf(chemin: str, verbose: bool = False) -> Optional[str]:
    """Extrait le texte d'un fichier PDF."""
    _check_pypdf()
    try:
        reader = PdfReader(chemin)
        pages = []
        for i, page in enumerate(reader.pages):
            try:
                texte_page = page.extract_text() or ""
                pages.append(texte_page)
            except Exception as e:
                if verbose:
                    print(f"    Page {i+1} : erreur extraction ({e})")
                continue
        if not pages:
            return None
        texte_complet = "\n\n".join(pages)
        return _nettoyer_texte(texte_complet)
    except Exception as e:
        if verbose:
            print(f"  Erreur PDF '{chemin}' : {e}")
        return None


# ---------------------------------------------------------------------------
# Extraction DOCX (texte + tableaux — crucial pour les BPU)
# ---------------------------------------------------------------------------

def extraire_docx(chemin: str, verbose: bool = False) -> Optional[str]:
    """
    Extrait le texte d'un fichier .docx (paragraphes + tableaux).

    Important : les BPU France Travail sont presque toujours des tableaux
    Word, donc on extrait DEUX choses :
    - les paragraphes (texte courant)
    - les cellules de tableaux (où sont les prix)

    Les tableaux sont sérialisés en format lisible pour le LLM :
        [Tableau 1]
        Libellé | Prix unitaire HT | Quantité estimée
        Parcours MILE complet | 1 200 € | 140
        ...
    """
    if not _docx_disponible():
        if verbose:
            print(f"  ⚠ python-docx non installé, .docx ignoré.")
            print(f"    Pour activer : pip install python-docx")
        return None

    try:
        doc = DocxDocument(chemin)
        blocs = []

        # 1. Paragraphes du corps
        for para in doc.paragraphs:
            texte = para.text.strip()
            if texte:
                blocs.append(texte)

        # 2. Tableaux — sérialisation lisible pour le LLM
        for i, table in enumerate(doc.tables):
            blocs.append(f"\n[Tableau {i+1}]")
            for row in table.rows:
                cellules = []
                for cell in row.cells:
                    # On nettoie chaque cellule (peut contenir des sauts de ligne)
                    txt = " ".join(cell.text.split())
                    cellules.append(txt)
                if any(c for c in cellules):  # ignore lignes vides
                    blocs.append(" | ".join(cellules))
            blocs.append("")  # ligne vide après chaque tableau

        # 3. En-têtes (parfois infos importantes)
        for section in doc.sections:
            for header in [section.header, section.first_page_header,
                            section.even_page_header]:
                if header is None:
                    continue
                for para in header.paragraphs:
                    txt = para.text.strip()
                    if txt:
                        blocs.append(f"[En-tête] {txt}")

        if not blocs:
            return None
        texte_complet = "\n".join(blocs)
        return _nettoyer_texte(texte_complet)
    except Exception as e:
        if verbose:
            print(f"  Erreur DOCX '{chemin}' : {e}")
        return None


# ---------------------------------------------------------------------------
# Extraction XLSX
# ---------------------------------------------------------------------------

def extraire_xlsx(chemin: str, verbose: bool = False) -> Optional[str]:
    """Extrait le contenu d'un fichier .xlsx (toutes feuilles, cellules non vides)."""
    if not _xlsx_disponible():
        if verbose:
            print(f"  ⚠ openpyxl non installé, .xlsx ignoré.")
            print(f"    Pour activer : pip install openpyxl")
        return None

    try:
        wb = openpyxl.load_workbook(chemin, data_only=True, read_only=True)
        blocs = []
        for nom_feuille in wb.sheetnames:
            ws = wb[nom_feuille]
            blocs.append(f"\n[Feuille: {nom_feuille}]")
            for row in ws.iter_rows(values_only=True):
                cellules = []
                for c in row:
                    if c is None:
                        cellules.append("")
                    elif isinstance(c, float) and c.is_integer():
                        cellules.append(str(int(c)))
                    else:
                        cellules.append(str(c).strip())
                if any(c for c in cellules):
                    blocs.append(" | ".join(cellules))
            blocs.append("")
        if not blocs:
            return None
        return _nettoyer_texte("\n".join(blocs))
    except Exception as e:
        if verbose:
            print(f"  Erreur XLSX '{chemin}' : {e}")
        return None


# ---------------------------------------------------------------------------
# Dispatcher par extension
# ---------------------------------------------------------------------------

def extraire_fichier(chemin: str, verbose: bool = False) -> Optional[str]:
    """
    Détecte l'extension et appelle l'extracteur approprié.
    Returns None si format non supporté ou échec.
    """
    ext = os.path.splitext(chemin)[1].lower()
    if ext == ".pdf":
        return extraire_pdf(chemin, verbose=verbose)
    if ext == ".docx":
        return extraire_docx(chemin, verbose=verbose)
    if ext == ".xlsx":
        return extraire_xlsx(chemin, verbose=verbose)
    return None


# ---------------------------------------------------------------------------
# Extraction d'un dossier complet
# ---------------------------------------------------------------------------

def extraire_dossier(dossier: str, verbose: bool = False) -> dict:
    """
    Extrait tous les fichiers supportés (.pdf, .docx, .xlsx) d'un dossier.

    Returns:
        {nom_fichier: {hash, texte, taille_chars, format}}
    """
    if not os.path.isdir(dossier):
        return {}
    resultats = {}

    # Pré-vérification des dépendances (warning unique en début)
    fichiers = os.listdir(dossier)
    a_docx = any(n.lower().endswith(".docx") for n in fichiers)
    a_xlsx = any(n.lower().endswith(".xlsx") for n in fichiers)
    if a_docx and not _docx_disponible() and verbose:
        print(f"  ⚠ Des fichiers .docx présents mais python-docx absent.")
        print(f"    pip install python-docx  (pour lire les BPU et cadres de réponse)")
    if a_xlsx and not _xlsx_disponible() and verbose:
        print(f"  ⚠ Des fichiers .xlsx présents mais openpyxl absent.")
        print(f"    pip install openpyxl  (pour lire les bases admin)")

    for nom in sorted(fichiers):
        ext = os.path.splitext(nom)[1].lower()
        if ext not in EXTENSIONS_SUPPORTEES:
            continue
        chemin = os.path.join(dossier, nom)
        if not os.path.isfile(chemin):
            continue
        h = hash_fichier(chemin)
        texte = extraire_fichier(chemin, verbose=verbose)
        if texte is None:
            if verbose:
                print(f"  ⚠ Impossible d'extraire {nom}")
            continue
        resultats[nom] = {
            "hash": h,
            "texte": texte,
            "taille_chars": len(texte),
            "format": ext.lstrip("."),
        }
        if verbose:
            print(f"  ✓ {nom} : {len(texte)} chars extraits ({ext})")
    return resultats


# ---------------------------------------------------------------------------
# CLI test
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage : python pdf_extract.py <chemin.pdf|.docx|.xlsx|dossier>")
        sys.exit(1)
    chemin = sys.argv[1]
    if os.path.isfile(chemin):
        texte = extraire_fichier(chemin, verbose=True)
        if texte:
            print(f"\n=== {len(texte)} caractères extraits ===")
            print(texte[:2000])
            if len(texte) > 2000:
                print(f"\n... ({len(texte)-2000} chars supplémentaires)")
    elif os.path.isdir(chemin):
        res = extraire_dossier(chemin, verbose=True)
        print(f"\n=== {len(res)} fichiers traités ===")
        for nom, info in res.items():
            print(f"  {nom} : {info['taille_chars']} chars, "
                  f"format {info['format']}, hash {info['hash'][:12]}")
